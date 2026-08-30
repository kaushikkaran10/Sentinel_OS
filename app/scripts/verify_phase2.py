"""Phase 2 functional verification — run, don't just import.

    cd app && ../.venv/Scripts/python.exe -m scripts.verify_phase2

Covers, per the approved plan:
  A. File generators   — docx / xlsx round-trip
  B. RAG               — offline model, chunking, Chroma CRUD, persistence
  D. Docker sandbox    — degraded string + live container (needs Docker running)

Section C (the /kb HTTP endpoints) is exercised separately with curl against a
running server. Prints PASS / FAIL / SKIP per check; exits non-zero on any FAIL.
"""

from __future__ import annotations

import sys
import traceback
from pathlib import Path

from core.logging import configure_logging

configure_logging()

_PASS = 0
_FAIL = 0
_SKIP = 0


def check(name: str, fn) -> None:
    global _PASS, _FAIL, _SKIP
    try:
        result = fn()
        if result == "SKIP":
            _SKIP += 1
            print(f"SKIP  {name}")
        else:
            _PASS += 1
            print(f"PASS  {name}")
    except Exception as exc:  # noqa: BLE001 - report everything
        _FAIL += 1
        print(f"FAIL  {name}: {type(exc).__name__}: {exc}")
        traceback.print_exc()


# ── A. File generators ─────────────────────────────────────────────────────
def _a_docx():
    from docx import Document

    from services.tools.file_maker import generate_approval_note

    p = Path(
        generate_approval_note(
            "Quarterly safety review complete.",
            {"defects_found": 3, "pass_rate": "98.2%", "inspector": "A. Rao"},
        )
    )
    assert p.exists() and p.suffix == ".docx", p
    assert p.parent.name == "generated", p
    assert p.stat().st_size > 0
    text = "\n".join(par.text for par in Document(str(p)).paragraphs)
    tables = Document(str(p)).tables
    cells = " ".join(c.text for t in tables for row in t.rows for c in row.cells)
    assert "Quarterly safety review complete." in text, text
    assert "defects_found" in cells and "98.2%" in cells, cells


def _a_xlsx():
    from openpyxl import load_workbook

    from services.tools.file_maker import generate_metrics_sheet

    p = Path(
        generate_metrics_sheet(
            [
                {"month": "Jan", "output": 120, "defects": 4},
                {"month": "Feb", "output": 135, "defects": 2},
            ]
        )
    )
    assert p.exists() and p.suffix == ".xlsx" and p.parent.name == "generated", p
    ws = load_workbook(str(p)).active
    assert [c.value for c in ws[1]] == ["month", "output", "defects"]
    assert ws.max_row == 3
    assert ws["B3"].value == 135, ws["B3"].value


def _a_xlsx_empty():
    from openpyxl import load_workbook

    from services.tools.file_maker import generate_metrics_sheet

    p = Path(generate_metrics_sheet([]))
    ws = load_workbook(str(p)).active
    assert ws["A1"].value == "No data"


# ── B. RAG ────────────────────────────────────────────────────────────────
_SOP = (
    "<html><body><h1>Welding Safety SOP</h1>"
    "<p>All welders must wear a full face visor, flame-resistant gloves, and a "
    "leather apron before striking an arc.</p>"
    "<p>Confirm the fume-extraction fan is running and the bay is clear of "
    "solvents.</p>"
    "<p>Argon shielding-gas cylinders must be chained upright at all times.</p>"
    "</body></html>"
) + (" Supplementary hot-work permit and ventilation guidance follows." * 80)


def _b_offline_env():
    import os

    import services.rag.vector_db as v  # noqa: F401

    assert os.environ.get("HF_HUB_OFFLINE") == "1"
    assert os.environ.get("TRANSFORMERS_OFFLINE") == "1"


def _b_model_local():
    import services.rag.vector_db as v

    m = v._get_model()
    dim = getattr(m, "get_embedding_dimension", m.get_sentence_embedding_dimension)()
    assert dim == 384, dim


def _b_chunking():
    import services.rag.vector_db as v

    chunks = v.chunk_text(_SOP)
    assert len(chunks) > 1, len(chunks)
    assert "<p>" not in chunks[0] and "</h1>" not in chunks[0]
    assert "  " not in chunks[0], "whitespace not normalized"
    # consecutive chunks must share an overlap span
    tail = chunks[0][-120:].split()
    assert any(w in chunks[1][:400] for w in tail if len(w) > 3), "no overlap detected"


def _b_crud():
    import services.rag.vector_db as v

    v.delete_document("verify-welding")  # idempotent cleanup
    n = v.add_document("verify-welding", "sop_welding.txt", "SOP", _SOP)
    assert n > 0, n
    docs = {d["id"]: d for d in v.list_documents()}
    assert docs["verify-welding"]["filename"] == "sop_welding.txt"
    assert docs["verify-welding"]["chunks"] == n

    ans = v.query_local_knowledge("What protective equipment is required for welding?")
    assert "[source: sop_welding.txt]" in ans, ans
    assert "visor" in ans.lower() or "gloves" in ans.lower(), ans

    before = v.collection_count()
    removed = v.delete_document("verify-welding")
    assert removed == n
    assert v.collection_count() == before - n
    assert "verify-welding" not in {d["id"] for d in v.list_documents()}


def _b_persistence():
    import chromadb

    import services.rag.vector_db as v
    from core.config import CHROMA_COLLECTION, CHROMA_DIR

    v.delete_document("verify-persist")
    v.add_document("verify-persist", "keep.txt", "SOP", "Torque spec for the mount bolt is 42 Nm.")
    assert (CHROMA_DIR / "chroma.sqlite3").exists(), "chroma sqlite not on disk"

    # brand-new client object, same on-disk path
    fresh = chromadb.PersistentClient(path=str(CHROMA_DIR))
    col = fresh.get_collection(CHROMA_COLLECTION)
    hits = col.get(where={"doc_id": "verify-persist"})
    assert len(hits["ids"]) >= 1, "doc did not persist to disk"
    v.delete_document("verify-persist")


def _b_embedded_mode():
    import inspect

    import services.rag.vector_db as v

    src = inspect.getsource(v)
    assert "PersistentClient" in src and "HttpClient" not in src, "must be embedded mode"
    assert isinstance(v._get_collection().name, str)


# ── D. Docker sandbox ─────────────────────────────────────────────────────
def _d_degraded():
    import services.tools.code_sandbox as cs
    from core.config import DOCKER_UNAVAILABLE_MSG
    from core.runtime import runtime

    called = {"v": False}
    orig = None
    try:
        import docker

        orig = docker.from_env

        def _boom(*a, **k):
            called["v"] = True
            raise AssertionError("docker.from_env must not be called in degraded mode")

        docker.from_env = _boom
        prev = runtime.docker_available
        runtime.docker_available = False
        out = cs.execute_sandbox_code("print('hello')")
        runtime.docker_available = prev
        assert out == DOCKER_UNAVAILABLE_MSG, repr(out)
        assert called["v"] is False
    finally:
        if orig is not None:
            docker.from_env = orig


def _d_live():
    import asyncio

    import services.tools.code_sandbox as cs
    from core.docker_health import check_docker
    from core.runtime import runtime

    ok, reason = asyncio.run(check_docker())
    runtime.docker_available = ok
    if not ok:
        print(f"      (Docker not reachable: {reason})")
        return "SKIP"

    out = cs.execute_sandbox_code("print('hello')")
    assert "hello" in out, repr(out)

    net = cs.execute_sandbox_code(
        "import socket\n"
        "try:\n"
        "    socket.create_connection(('1.1.1.1', 53), 2); print('NET_OK')\n"
        "except Exception as e:\n"
        "    print('NET_BLOCKED', type(e).__name__)"
    )
    assert "NET_BLOCKED" in net, repr(net)

    err = cs.execute_sandbox_code("import sys; sys.stderr.write('boom')")
    assert "boom" in err, repr(err)

    slow = cs.execute_sandbox_code("while True:\n    pass")
    assert "timed out" in slow, repr(slow)

    import subprocess

    leftovers = subprocess.run(
        ["docker", "ps", "-a", "--filter", "name=sentinel-sandbox-", "--format", "{{.Names}}"],
        capture_output=True, text=True,
    ).stdout.strip()
    assert not leftovers, f"leftover containers: {leftovers}"


def main() -> int:
    print("=== A. File generators ===")
    check("docx approval note round-trips", _a_docx)
    check("xlsx metrics sheet round-trips", _a_xlsx)
    check("xlsx empty input is still valid", _a_xlsx_empty)

    print("\n=== B. RAG (ChromaDB + offline embeddings) ===")
    check("HF offline env vars forced", _b_offline_env)
    check("embedding model loads from local disk", _b_model_local)
    check("chunking: tags stripped, overlap present", _b_chunking)
    check("Chroma add / list / query / delete", _b_crud)
    check("PersistentClient writes survive a fresh client", _b_persistence)
    check("embedded mode, no HttpClient", _b_embedded_mode)

    print("\n=== D. Docker sandbox ===")
    check("degraded mode returns the exact string, no SDK call", _d_degraded)
    check("live container: hello / no-net / stderr / timeout / cleanup", _d_live)

    print(f"\n{'=' * 44}\nPASS={_PASS}  FAIL={_FAIL}  SKIP={_SKIP}")
    return 1 if _FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
