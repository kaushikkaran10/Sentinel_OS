"""Local deliverable generators (4_Agent_Logic_&_Tools.md §3).

Two of the four fixed MVP tools:

* ``generate_approval_note(summary, metrics) -> str`` — a formatted Word document.
* ``generate_metrics_sheet(data) -> str``           — an Excel workbook.

Both write into ``data/generated/`` and return the **absolute path** of the file
they created. Signatures are kept synchronous exactly as the spec defines them;
these are blocking (python-docx / openpyxl do real file IO), so the Phase 4 Tool
Execution Node will call them via ``run_in_threadpool`` rather than awaiting.
"""

from __future__ import annotations

import uuid
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.config import GENERATED_DIR
from core.logging import get_logger

logger = get_logger("sentinel.tools.file_maker")

_MAX_COL_WIDTH = 60  # openpyxl column-width units, keeps wide cells sane


def _unique_path(prefix: str, ext: str) -> Path:
    """``data/generated/<prefix>_<8 hex>.<ext>`` — collision-safe, no user input."""
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    return GENERATED_DIR / f"{prefix}_{uuid.uuid4().hex[:8]}.{ext}"


def _timestamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def generate_approval_note(summary: str, metrics: dict) -> str:
    """Create a formatted Word approval note and return its absolute path.

    ``metrics`` is rendered as a two-column key/value table. An empty mapping is
    valid — the table is replaced with a short note.
    """
    doc = Document()

    doc.add_heading("Approval Note", level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Generated: {_timestamp()}").italic = True

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(str(summary).strip() or "(no summary provided)")

    doc.add_heading("Metrics", level=1)
    if metrics:
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text = "Metric", "Value"
        for cell in hdr:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.bold = True
                    run.font.size = Pt(10)
        for key, value in metrics.items():
            row = table.add_row().cells
            row[0].text = str(key)
            row[1].text = str(value)
    else:
        doc.add_paragraph("No metrics provided.")

    path = _unique_path("approval_note", "docx")
    doc.save(str(path))
    logger.info("Generated approval note: %s", path.name)
    return str(path)


def generate_metrics_sheet(data: list[dict]) -> str:
    """Create an Excel metrics sheet from a list of row dicts; return its path.

    The header is the ordered union of keys across every row (first-seen order).
    Missing keys become blank cells. An empty list still yields a valid workbook.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Metrics"

    columns: list[str] = []
    for row in data:
        for key in row:
            if key not in columns:
                columns.append(key)

    if not columns:
        ws["A1"] = "No data"
        path = _unique_path("metrics_sheet", "xlsx")
        wb.save(str(path))
        logger.info("Generated metrics sheet (empty): %s", path.name)
        return str(path)

    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for row in data:
        ws.append([_cell(row.get(col, "")) for col in columns])

    _autosize(ws, columns, data)

    path = _unique_path("metrics_sheet", "xlsx")
    wb.save(str(path))
    logger.info("Generated metrics sheet: %s rows, %s cols -> %s", len(data), len(columns), path.name)
    return str(path)


def _cell(value: object) -> object:
    """Keep native numbers/bools; stringify everything else so openpyxl is happy."""
    if isinstance(value, (int, float, bool)) or value is None:
        return value
    return str(value)


def _autosize(ws, columns: list[str], data: list[dict]) -> None:
    for idx, col in enumerate(columns, start=1):
        longest = len(str(col))
        for row in data:
            longest = max(longest, len(str(row.get(col, ""))))
        ws.column_dimensions[get_column_letter(idx)].width = min(longest + 2, _MAX_COL_WIDTH)


if __name__ == "__main__":  # pragma: no cover - manual smoke test
    note = generate_approval_note(
        "Quarterly safety review complete. All critical findings closed.",
        {"defects_found": 3, "pass_rate": "98.2%", "inspector": "A. Rao"},
    )
    sheet = generate_metrics_sheet(
        [
            {"month": "Jan", "output": 120, "defects": 4},
            {"month": "Feb", "output": 135, "defects": 2},
        ]
    )
    print("approval note:", note)
    print("metrics sheet:", sheet)
